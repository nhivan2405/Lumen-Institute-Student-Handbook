"""Đọc bốn định dạng được chấp nhận và luôn giữ metadata nguồn thật."""
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt"}

def extract_document(path: Path) -> list[dict]:
    """Trả các đơn vị text `{text, location}`; PDF/page và PPTX/slide không bị bịa."""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("Chỉ hỗ trợ PDF, DOCX, PPTX và TXT.")
    if suffix == ".pdf":
        import fitz
        with fitz.open(path) as pdf:
            return [{"text": page.get_text("text"), "location": f"trang {index + 1}"}
                    for index, page in enumerate(pdf) if page.get_text("text").strip()]
    if suffix == ".docx":
        from docx import Document
        document = Document(path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
        return [{"text": text, "location": "nội dung DOCX"}] if text.strip() else []
    if suffix == ".pptx":
        from pptx import Presentation
        presentation = Presentation(path)
        result = []
        for index, slide in enumerate(presentation.slides):
            texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            if texts:
                result.append({"text": "\n".join(texts), "location": f"slide {index + 1}"})
        return result
    return [{"text": path.read_text(encoding="utf-8", errors="replace"), "location": "nội dung TXT"}]
